import csv
from datetime import datetime

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from database import Project

HEADER_BG = colors.HexColor("#2E4053")
GRID_COLOR = colors.HexColor("#CCCCCC")


def _styled_table(data: list[list[str]]) -> Table:
    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black),
                ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.25, GRID_COLOR),
            ]
        )
    )
    return table


def export_pdf(project: Project, filepath: str) -> None:
    doc = SimpleDocTemplate(filepath, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"Cost Report — {project.name}", styles["Title"]),
        Paragraph(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Spacer(1, 0.25 * inch),
        Paragraph("Materials", styles["Heading2"]),
    ]

    if project.materials:
        data = [["Dr. No", "Date", "Supplier", "Item", "Qty", "Total"]]
        for m in project.materials:
            data.append([m.dr_no, m.date, m.supplier, m.name, f"{m.qty_number:g} {m.qty_unit}".strip(),
                         f"${m.total:,.2f}"])
        data.append(["", "", "", "", "Materials Total", f"${project.materials_total:,.2f}"])
        elements.append(_styled_table(data))
    else:
        elements.append(Paragraph("No material line items.", styles["Normal"]))
    elements.append(Spacer(1, 0.25 * inch))

    elements.append(Paragraph("Payroll", styles["Heading2"]))
    if project.labor:
        data = [["Payroll No.", "Date Range", "Worker", "Skill", "Days", "Rate",
                 "Cash Advance", "OT Hours", "OT Pay", "Total"]]
        for l in project.labor:
            date_range = f"{l.date_start} to {l.date_end}" if l.date_start or l.date_end else ""
            data.append([l.payroll_no, date_range, l.worker, l.skill, f"{l.days:g}",
                         f"${l.rate:,.2f}", f"${l.cash_advance:,.2f}", f"{l.overtime_hours:g}",
                         f"${l.overtime_pay:,.2f}", f"${l.total:,.2f}"])
        data.append(["", "", "", "", "", "", "", "", "Labor Total", f"${project.labor_total:,.2f}"])
        elements.append(_styled_table(data))
    else:
        elements.append(Paragraph("No labor line items.", styles["Normal"]))
    elements.append(Spacer(1, 0.25 * inch))

    elements.append(Paragraph("Equipment", styles["Heading2"]))
    if project.equipment:
        data = [["Equipment", "Days", "Rate", "Total"]]
        for eq in project.equipment:
            data.append([eq.equipment, f"{eq.days:g}", f"${eq.rate:,.2f}", f"${eq.total:,.2f}"])
        data.append(["", "", "Equipment Total", f"${project.equipment_total:,.2f}"])
        elements.append(_styled_table(data))
    else:
        elements.append(Paragraph("No equipment line items.", styles["Normal"]))
    elements.append(Spacer(1, 0.25 * inch))

    elements.append(Paragraph("Expenses", styles["Heading2"]))
    if project.expenses:
        data = [["Receipt No.", "Date", "Paid To", "Description", "Amount"]]
        for x in project.expenses:
            data.append([x.receipt_no, x.date, x.paid_to, x.description, f"${x.amount:,.2f}"])
        data.append(["", "", "", "Expenses Total", f"${project.expenses_total:,.2f}"])
        elements.append(_styled_table(data))
    else:
        elements.append(Paragraph("No expense line items.", styles["Normal"]))
    elements.append(Spacer(1, 0.25 * inch))

    elements.append(Paragraph("Admin Expenses", styles["Heading2"]))
    if project.admin_expenses:
        data = [["Receipt No.", "Date", "Paid To", "Description", "Amount"]]
        for x in project.admin_expenses:
            data.append([x.receipt_no, x.date, x.paid_to, x.description, f"${x.amount:,.2f}"])
        data.append(["", "", "", "Admin Expenses Total", f"${project.admin_expenses_total:,.2f}"])
        elements.append(_styled_table(data))
    else:
        elements.append(Paragraph("No admin expense line items.", styles["Normal"]))
    elements.append(Spacer(1, 0.35 * inch))

    elements.append(Paragraph(f"Project Cost: ${project.total_cost:,.2f}", styles["Heading1"]))

    doc.build(elements)


def export_docx(project: Project, filepath: str) -> None:
    doc = Document()
    doc.add_heading(f"Cost Report — {project.name}", level=0)
    doc.add_paragraph(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    doc.add_heading("Materials", level=1)
    if project.materials:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ["Dr. No", "Date", "Supplier", "Item", "Qty", "Total"]):
            cell.text = text
        for m in project.materials:
            cells = table.add_row().cells
            cells[0].text = m.dr_no
            cells[1].text = m.date
            cells[2].text = m.supplier
            cells[3].text = m.name
            cells[4].text = f"{m.qty_number:g} {m.qty_unit}".strip()
            cells[5].text = f"${m.total:,.2f}"
        doc.add_paragraph(f"Materials Total: ${project.materials_total:,.2f}").runs[0].bold = True
    else:
        doc.add_paragraph("No material line items.")

    doc.add_heading("Payroll", level=1)
    if project.labor:
        table = doc.add_table(rows=1, cols=10)
        table.style = "Light Grid Accent 1"
        headers = ["Payroll No.", "Date Range", "Worker", "Skill", "Days", "Rate",
                   "Cash Advance", "OT Hours", "OT Pay", "Total"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for l in project.labor:
            cells = table.add_row().cells
            cells[0].text = l.payroll_no
            cells[1].text = f"{l.date_start} to {l.date_end}" if l.date_start or l.date_end else ""
            cells[2].text = l.worker
            cells[3].text = l.skill
            cells[4].text = f"{l.days:g}"
            cells[5].text = f"${l.rate:,.2f}"
            cells[6].text = f"${l.cash_advance:,.2f}"
            cells[7].text = f"{l.overtime_hours:g}"
            cells[8].text = f"${l.overtime_pay:,.2f}"
            cells[9].text = f"${l.total:,.2f}"
        doc.add_paragraph(f"Labor Total: ${project.labor_total:,.2f}").runs[0].bold = True
    else:
        doc.add_paragraph("No labor line items.")

    doc.add_heading("Equipment", level=1)
    if project.equipment:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ["Equipment", "Days", "Rate", "Total"]):
            cell.text = text
        for eq in project.equipment:
            cells = table.add_row().cells
            cells[0].text = eq.equipment
            cells[1].text = f"{eq.days:g}"
            cells[2].text = f"${eq.rate:,.2f}"
            cells[3].text = f"${eq.total:,.2f}"
        doc.add_paragraph(f"Equipment Total: ${project.equipment_total:,.2f}").runs[0].bold = True
    else:
        doc.add_paragraph("No equipment line items.")

    doc.add_heading("Expenses", level=1)
    if project.expenses:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ["Receipt No.", "Date", "Paid To", "Description", "Amount"]):
            cell.text = text
        for x in project.expenses:
            cells = table.add_row().cells
            cells[0].text = x.receipt_no
            cells[1].text = x.date
            cells[2].text = x.paid_to
            cells[3].text = x.description
            cells[4].text = f"${x.amount:,.2f}"
        doc.add_paragraph(f"Expenses Total: ${project.expenses_total:,.2f}").runs[0].bold = True
    else:
        doc.add_paragraph("No expense line items.")

    doc.add_heading("Admin Expenses", level=1)
    if project.admin_expenses:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ["Receipt No.", "Date", "Paid To", "Description", "Amount"]):
            cell.text = text
        for x in project.admin_expenses:
            cells = table.add_row().cells
            cells[0].text = x.receipt_no
            cells[1].text = x.date
            cells[2].text = x.paid_to
            cells[3].text = x.description
            cells[4].text = f"${x.amount:,.2f}"
        doc.add_paragraph(f"Admin Expenses Total: ${project.admin_expenses_total:,.2f}").runs[0].bold = True
    else:
        doc.add_paragraph("No admin expense line items.")

    total_paragraph = doc.add_paragraph()
    run = total_paragraph.add_run(f"Project Cost: ${project.total_cost:,.2f}")
    run.bold = True
    run.font.size = Pt(16)

    doc.save(filepath)


def export_csv(project: Project, filepath: str) -> None:
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([f"Cost Report - {project.name}"])
        writer.writerow([f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
        writer.writerow([])

        writer.writerow(["Materials"])
        writer.writerow(["Dr. No", "Date", "Supplier", "Item", "Qty Number", "Qty Unit", "Total"])
        for m in project.materials:
            writer.writerow([m.dr_no, m.date, m.supplier, m.name, m.qty_number, m.qty_unit, m.total])
        writer.writerow(["", "", "", "", "", "Materials Total", project.materials_total])
        writer.writerow([])

        writer.writerow(["Payroll"])
        writer.writerow(["Payroll No.", "Date Start", "Date End", "Worker", "Skill", "Days Marked",
                         "Days", "Rate", "Cash Advance", "OT Hours", "OT Pay", "Total"])
        for l in project.labor:
            writer.writerow([l.payroll_no, l.date_start, l.date_end, l.worker, l.skill,
                             " / ".join(l.days_marked), l.days, l.rate, l.cash_advance,
                             l.overtime_hours, l.overtime_pay, l.total])
        writer.writerow(["", "", "", "", "", "", "", "", "", "", "Labor Total", project.labor_total])
        writer.writerow([])

        writer.writerow(["Equipment"])
        writer.writerow(["Equipment", "Days", "Rate", "Total"])
        for eq in project.equipment:
            writer.writerow([eq.equipment, eq.days, eq.rate, eq.total])
        writer.writerow(["", "", "Equipment Total", project.equipment_total])
        writer.writerow([])

        writer.writerow(["Expenses"])
        writer.writerow(["Receipt No.", "Date", "Paid To", "Description", "Amount"])
        for x in project.expenses:
            writer.writerow([x.receipt_no, x.date, x.paid_to, x.description, x.amount])
        writer.writerow(["", "", "", "Expenses Total", project.expenses_total])
        writer.writerow([])

        writer.writerow(["Admin Expenses"])
        writer.writerow(["Receipt No.", "Date", "Paid To", "Description", "Amount"])
        for x in project.admin_expenses:
            writer.writerow([x.receipt_no, x.date, x.paid_to, x.description, x.amount])
        writer.writerow(["", "", "", "Admin Expenses Total", project.admin_expenses_total])
        writer.writerow([])

        writer.writerow(["Project Cost", project.total_cost])


def export_projects_summary_csv(projects: list[Project], filepath: str) -> None:
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Project", "Materials Total", "Labor Total", "Equipment Total", "Expenses Total",
                         "Admin Expenses Total", "Project Cost", "Material Items", "Labor Items",
                         "Equipment Items", "Expense Items", "Admin Expense Items"])
        for p in projects:
            writer.writerow([p.name, p.materials_total, p.labor_total, p.equipment_total, p.expenses_total,
                             p.admin_expenses_total, p.total_cost, len(p.materials), len(p.labor),
                             len(p.equipment), len(p.expenses), len(p.admin_expenses)])
